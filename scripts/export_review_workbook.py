"""검토용 엑셀: 영화관 전수 목록 · 연도별 유형 · 개업·폐업·인수합병 연표 · 연도별 집계 · 검수 결과 · 보정 내역.

원자료를 모두 담은 data/seoul_cinema_master.xlsx 와 달리, 사람이 읽고 확인하기 쉽게 한국어 열 이름으로 정리한다.
산출: data/seoul_cinema_dataset.xlsx
"""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
CLEAN = ROOT / "data" / "clean"
OUT = ROOT / "data" / "seoul_cinema_dataset.xlsx"
TYPE_FILL = {"멀티플렉스": "DCE3F2", "독립·예술영화전용관": "F9DDD6", "일반영화관": "D4EEEA"}
SEV_FILL = {"오류": "F6C9C4", "확인 필요": "FBE8B5", "참고": "EEEEEE"}

GUIDE = [
    ("이 파일", "서울 영화관 연구 자료(검토용). LOCALDATA(행정안전부 영화상영관 인허가) 1차 기준 + 영화진흥위원회(KOFIC) 극장 현황 2010~2025 보완"),
    ("기준일", "LOCALDATA 2026-09-04 갱신분 · KOFIC 2025년 말"),
    ("분석 단위", "운영구간 = 한 장소에서 인허가가 이어진 개업~폐업 한 번. 같은 자리에서 새 인허가로 다시 열면 다른 운영구간"),
    ("유형(영화진흥위원회 구분)", "멀티플렉스 체인(4대 체인 CGV·롯데시네마·메가박스·씨네Q 브랜드 운영, 직영·위탁) · 독립·예술영화전용관 · 일반영화관(비체인). 작은영화관은 서울에 없어 제외"),
    ("유형은 해마다 다를 수 있음", "일반영화관이 체인에 인수되거나 체인 브랜드가 바뀌면 그해부터 유형·브랜드가 바뀜 → '연도별_유형' 시트. '영화관_전수'의 '현재(마지막) 유형'은 영업 중이면 지금, 폐업이면 문 닫을 때 유형"),
    ("영화관_전수", "운영구간 전체 목록(영업·폐업 모두). 유형 변천 · 직영/위탁 · 날짜 · 위치 · 규모 · 입점 형태"),
    ("연도별_유형", "운영구간 × 연도의 유형·브랜드·판단 근거(KOFIC 그해 기록 / 앞뒤 해 값 / LOCALDATA 상호 / 확인값)"),
    ("사건연표", "개업 · 폐업 · 체인 인수·가맹 · 체인 이탈 · 브랜드 변경(인수합병) · 같은 자리 승계, 연도순"),
    ("연도별_집계", "유형별 운영 수 · 개업 · 폐업 · 인수·전환 들어옴/나감 · 스크린(KOFIC 수록분). 유형은 그해 유형"),
    ("검수결과", "규칙별 자동 검수: 오류(고쳐야 함) · 확인 필요(자료끼리 다름) · 참고(한계)"),
    ("보정내역", "다른 자료로 확인해 고친 값(상태·폐업일·좌표·상호)과 유형 확인값"),
    ("주의", "2010년 전(KOFIC 기록 없음)의 인수 연도는 KOFIC 개관일·LOCALDATA 상호로 추정. LOCALDATA 개업일 중 일괄 등록일(1998-03-09 강남, 2003-02-12 종로, 2003-05-20 중구)은 실제 개관일이 아님"),
]


def main():
    e = pd.read_csv(CLEAN / "venues_episodes.csv", dtype=str).fillna("")
    import json
    founded = {f["properties"]["id"]: f["properties"].get("founded") for f in json.loads((ROOT / "story" / "public" / "data" / "venues.geojson").read_text(encoding="utf-8"))["features"]}
    floor = {f["properties"]["id"]: f["properties"].get("floor") for f in json.loads((ROOT / "story" / "public" / "data" / "venues.geojson").read_text(encoding="utf-8"))["features"]}
    main_t = pd.DataFrame({
        "운영구간 ID": e.operation_episode_id, "장소 ID": e.venue_id, "영화관": e.standardized_name,
        "현재(마지막) 유형": e.analysis_type, "개업 때 유형": e.type_at_open, "유형 변천": e.type_history,
        "체인 운영(직영·위탁)": e.chain_operation, "세부 구분": e.facility_subtype, "상태": e.operation_status,
        "개업일(LOCALDATA)": e.open_date, "실제 개관 연도(참고 자료)": e.operation_episode_id.map(founded).fillna(""),
        "폐업일": e.close_date, "자치구": e.gu, "행정동": e.dong, "도로명주소": e.address_road, "지번주소": e.address_jibun,
        "경도": e.lon, "위도": e.lat, "영화관 층(주소)": e.operation_episode_id.map(floor).fillna(""),
        "스크린(KOFIC 최근)": e.screens_kofic_latest, "좌석(KOFIC 최근)": e.seats_kofic_latest, "스크린 2025": e.screens_kofic_2025, "좌석 2025": e.seats_kofic_2025,
        "KOFIC 수록 연도": e.kofic_years, "입점 형태": e.building_occupancy_type, "건물": e.building_name, "입점 형태 확신도": e.occupancy_confidence,
        "LOCALDATA 상호(인허가별)": e.source_name_localdata, "KOFIC 상호 이력": e.source_name_kofic,
        "유형 근거": e.type_basis, "유형 메모": e.type_note, "인허가 수": e.license_count, "자료 충돌": e.source_conflict, "검증·보정 메모": e.verification_note,
    })
    ty = pd.read_csv(CLEAN / "episode_type_year.csv", dtype=str).fillna("")
    names = e.set_index("operation_episode_id").standardized_name
    ty_t = pd.DataFrame({"운영구간 ID": ty.operation_episode_id, "영화관": ty.operation_episode_id.map(names), "연도": ty.year.astype(int), "유형": ty.type,
                         "브랜드": ty.brand, "직영·위탁": ty.chain_operation, "그해 KOFIC 상호": ty.kofic_name, "판단 근거": ty.basis})
    ev = pd.read_csv(CLEAN / "cinema_events_timeline.csv", dtype=str).fillna("")
    ev_t = pd.DataFrame({"연도": ev.year.astype(int), "날짜": ev.date, "사건": ev.event, "영화관": ev.name, "자치구": ev.gu,
                         "이전 유형": ev.type_before, "이전 브랜드": ev.brand_before, "이후 유형": ev.type_after, "이후 브랜드": ev.brand_after,
                         "내용": ev.note, "운영구간 ID": ev.operation_episode_id, "장소 ID": ev.venue_id})
    y = pd.read_csv(CLEAN / "yearly_series.csv")
    y_t = y.rename(columns={"type": "유형", "year": "연도", "partial": "부분 연도", "operating": "운영 중", "opened": "개업", "opened_unsure": "개업(일괄 등록일·불확실)",
                            "closed": "폐업", "conv_in": "인수·전환 들어옴", "conv_out": "인수·전환 나감", "cum_opened": "누적 개업", "cum_closed": "누적 폐업",
                            "screens": "스크린(KOFIC)", "seats": "좌석(KOFIC)", "special": "특별관 스크린", "kofic_n": "KOFIC 수록 극장 수"})
    rv = pd.read_csv(ROOT / "data" / "verify" / "dataset_review.csv", dtype=str).fillna("")
    corr = pd.read_csv(ROOT / "data" / "manual" / "episode_corrections.csv", dtype=str).fillna("").rename(columns={"standardized_name": "영화관", "field": "항목", "value": "고친 값", "source": "근거", "note": "메모"})
    th = pd.read_csv(ROOT / "data" / "manual" / "type_history.csv", dtype=str, comment="#").fillna("")
    corr = pd.concat([corr, pd.DataFrame({"영화관": th.standardized_name, "항목": "유형(" + th.from_year + "년부터)", "고친 값": th.type, "근거": th.basis, "메모": "유형 확인값"})], ignore_index=True)

    with pd.ExcelWriter(OUT, engine="openpyxl") as w:
        pd.DataFrame(GUIDE, columns=["항목", "설명"]).to_excel(w, sheet_name="안내", index=False)
        main_t.to_excel(w, sheet_name="영화관_전수", index=False)
        ty_t.to_excel(w, sheet_name="연도별_유형", index=False)
        ev_t.to_excel(w, sheet_name="사건연표", index=False)
        y_t.to_excel(w, sheet_name="연도별_집계", index=False)
        rv.to_excel(w, sheet_name="검수결과", index=False)
        corr.to_excel(w, sheet_name="보정내역", index=False)
        head = PatternFill("solid", fgColor="2B2F36")
        for ws in w.book.worksheets:
            ws.freeze_panes = "B2" if ws.title in ("영화관_전수", "연도별_유형") else "A2"
            ws.auto_filter.ref = ws.dimensions
            for c in ws[1]:
                c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), head, Alignment(vertical="center", wrap_text=True)
            ws.row_dimensions[1].height = 30
            for col in ws.columns:
                width = max(len(str(c.value or "")) for c in list(col)[:300])
                ws.column_dimensions[col[0].column_letter].width = min(max(8, width * 1.15), 48 if ws.title != "안내" else 110)
            hdr = [c.value for c in ws[1]]
            for name in ("현재(마지막) 유형", "개업 때 유형", "유형", "이후 유형", "이전 유형"):
                if name in hdr:
                    j = hdr.index(name) + 1
                    for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                        if row[0].value in TYPE_FILL:
                            row[0].fill = PatternFill("solid", fgColor=TYPE_FILL[row[0].value])
            if "심각도" in hdr:
                j = hdr.index("심각도") + 1
                for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    if row[0].value in SEV_FILL:
                        row[0].fill = PatternFill("solid", fgColor=SEV_FILL[row[0].value])
        w.book["안내"].column_dimensions["A"].width = 24
        for row in w.book["안내"].iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
    print(f"→ {OUT.relative_to(ROOT)}  영화관 {len(main_t)} · 연도별 유형 {len(ty_t)} · 사건 {len(ev_t)} · 검수 {len(rv)}")


if __name__ == "__main__":
    main()
